from telebot import types
from loader import bot, db, marzban
from config import Config
from utils import keyboards
import os
import tempfile
import openpyxl
import openpyxl.utils
from datetime import datetime

# --- КОМАНДА /admin ---
@bot.message_handler(commands=['admin'])
def admin_start(message):
    if message.from_user.id == Config.ADMIN_ID:
        current_price = int(db.get_setting("sub_price", "150"))
        trial_days = int(db.get_setting("trial_days", "0"))
        ref_percent = int(db.get_setting("ref_percent", "20"))
        min_topup = int(db.get_setting("min_topup", "10"))
        bot.send_message(message.chat.id, "👨‍💻 Админ-панель:", reply_markup=keyboards.admin_menu(current_price, trial_days, ref_percent, min_topup))
    else:
        # Если пишет не админ — игнорим или отвечаем стандартно
        pass

# ==========================================
#          НАСТРОЙКИ БОТА
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_sales_menu")
def admin_sales_menu(call):
    """Открывает меню управления продажами"""
    if call.from_user.id != Config.ADMIN_ID: return
    bot.edit_message_text(
        "🔧 <b>Управление продажами</b>\n\nНажмите на пункт, чтобы вкл/выкл:",
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.sales_control_menu(db),
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda c: c.data.startswith("toggle_sell_"))
def toggle_sale_setting(call):
    """Переключает отдельную настройку продаж"""
    if call.from_user.id != Config.ADMIN_ID: return
    
    # Извлекаем ключ, например "sell_subs" из "toggle_sell_subs"
    key = call.data[len("toggle_"):]
    current = db.get_setting(key, "1")
    new_val = "0" if current == "1" else "1"
    db.set_setting(key, new_val)
    
    label = keyboards.SALES_TOGGLES.get(key, key)
    status = "включено" if new_val == "1" else "выключено"
    bot.answer_callback_query(call.id, f"{label} — {status}")
    
    # Обновляем меню
    bot.edit_message_reply_markup(
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.sales_control_menu(db)
    )

@bot.callback_query_handler(func=lambda c: c.data == "admin_back")
def admin_back(call):
    """Возврат в главное админ-меню"""
    if call.from_user.id != Config.ADMIN_ID: return
    current_price = int(db.get_setting("sub_price", "150"))
    trial_days = int(db.get_setting("trial_days", "0"))
    ref_percent = int(db.get_setting("ref_percent", "20"))
    bot.edit_message_text(
        "👨‍💻 Админ-панель:",
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.admin_menu(current_price, trial_days, ref_percent)
    )


@bot.callback_query_handler(func=lambda c: c.data == "admin_change_trial")
def change_trial_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите количество дней бесплатного триала (0 = отключен):")
    bot.register_next_step_handler(msg, process_change_trial, call.message.message_id)

def process_change_trial(message, admin_msg_id):
    if not message.text.isdigit():
        bot.send_message(message.chat.id, "Ошибка! Количество дней должно быть числом.")
        return
        
    new_trial = int(message.text)
    db.set_setting("trial_days", str(new_trial))
    
    status = f"{new_trial} дн." if new_trial > 0 else "отключен"
    bot.send_message(message.chat.id, f"✅ Бесплатный триал: {status}")
    
    # Обновляем админ-панель
    current_price = int(db.get_setting("sub_price", "150"))
    ref_percent = int(db.get_setting("ref_percent", "20"))
    min_topup = int(db.get_setting("min_topup", "10"))
    try:
        bot.edit_message_reply_markup(
            chat_id=message.chat.id, 
            message_id=admin_msg_id, 
            reply_markup=keyboards.admin_menu(current_price, new_trial, ref_percent, min_topup)
        )
    except Exception:
        pass

@bot.callback_query_handler(func=lambda c: c.data == "admin_change_ref_percent")
def change_ref_percent_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите процент реферального вознаграждения (от пополнений):")
    bot.register_next_step_handler(msg, process_change_ref_percent, call.message.message_id)

def process_change_ref_percent(message, admin_msg_id):
    if not message.text.isdigit():
        bot.send_message(message.chat.id, "Ошибка! Процент должен быть числом.")
        return
        
    new_percent = int(message.text)
    db.set_setting("ref_percent", str(new_percent))
    
    bot.send_message(message.chat.id, f"✅ Реферальный процент изменён на {new_percent}%!")
    
    # Обновляем админ-панель
    current_price = int(db.get_setting("sub_price", "150"))
    trial_days = int(db.get_setting("trial_days", "0"))
    min_topup = int(db.get_setting("min_topup", "10"))
    try:
        bot.edit_message_reply_markup(
            chat_id=message.chat.id, 
            message_id=admin_msg_id, 
            reply_markup=keyboards.admin_menu(current_price, trial_days, new_percent, min_topup)
        )
    except Exception:
        pass

@bot.callback_query_handler(func=lambda c: c.data == "admin_change_min_topup")
def change_min_topup_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите минимальную сумму пополнения (в рублях):")
    bot.register_next_step_handler(msg, process_change_min_topup, call.message.message_id)

def process_change_min_topup(message, admin_msg_id):
    if not message.text or not message.text.isdigit():
        bot.send_message(message.chat.id, "Ошибка! Сумма должна быть числом.")
        return

    new_min = int(message.text)
    if new_min < 1:
        bot.send_message(message.chat.id, "Ошибка! Минимум 1 рубль.")
        return

    db.set_setting("min_topup", str(new_min))
    bot.send_message(message.chat.id, f"✅ Минимальное пополнение изменено на {new_min}₽!")

    # Обновляем админ-панель
    current_price = int(db.get_setting("sub_price", "150"))
    trial_days = int(db.get_setting("trial_days", "0"))
    ref_percent = int(db.get_setting("ref_percent", "20"))
    try:
        bot.edit_message_reply_markup(
            chat_id=message.chat.id,
            message_id=admin_msg_id,
            reply_markup=keyboards.admin_menu(current_price, trial_days, ref_percent, new_min)
        )
    except Exception:
        pass

# ==========================================
#          ВЫДАЧА БАЛАНСА
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_give_balance")
def give_balance_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите ID пользователя:")
    bot.register_next_step_handler(msg, get_user_id_for_balance)

def get_user_id_for_balance(message):
    try:
        user_id = int(message.text)
        # Сохраняем ID и спрашиваем сумму
        msg = bot.send_message(message.chat.id, f"Введите сумму для пополнения ID {user_id}:")
        bot.register_next_step_handler(msg, lambda m: give_balance_final(m, user_id))
    except Exception:
        bot.send_message(message.chat.id, "❌ Это не ID (должно быть число).")

def give_balance_final(message, user_id):
    try:
        amount = int(message.text)
        db.update_balance(user_id, amount)
        bot.send_message(message.chat.id, f"✅ Баланс пользователя {user_id} пополнен на {amount}₽")
        # Пытаемся уведомить юзера
        try:
            bot.send_message(user_id, f"🎁 Администратор пополнил ваш баланс на {amount}₽")
        except Exception: pass
    except Exception:
        bot.send_message(message.chat.id, "❌ Неверная сумма.")

# ==========================================
#          ВЫДАЧА ПОДПИСКИ
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_grant_sub")
def grant_sub_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите ID пользователя для выдачи подписки:")
    bot.register_next_step_handler(msg, grant_sub_get_user)

def grant_sub_get_user(message):
    try:
        user_id = int(message.text)
        user = db.get_user(user_id)
        if not user:
            bot.send_message(message.chat.id, "❌ Пользователь не найден.")
            return
        bot.send_message(
            message.chat.id,
            f"Выберите срок подписки для ID {user_id}:",
            reply_markup=keyboards.admin_grant_sub_menu(user_id)
        )
    except ValueError:
        bot.send_message(message.chat.id, "❌ Это не ID (должно быть число).")

@bot.callback_query_handler(func=lambda c: c.data.startswith("agrant_"))
def grant_sub_final(call):
    """ callback: agrant_{user_id}_{days} """
    if call.from_user.id != Config.ADMIN_ID: return
    parts = call.data.split("_")
    user_id = int(parts[1])
    days = int(parts[2])

    # Обработка тестовой подписки на 1 минуту
    if days == -3:
        cluster = db.get_user_cluster(user_id)
        if not cluster:
            cluster = db.allocate_cluster()
            if not cluster:
                bot.edit_message_text(
                    "❌ Нет свободных мест на серверах для выдачи тестовой подписки.",
                    call.message.chat.id, call.message.message_id
                )
                return
            db.set_user_cluster(user_id, cluster)
            db.change_cluster_users_count(cluster, 1)

        import time, requests
        now = int(time.time())
        expire = now + 60
        
        # Обновляем БД
        with db.conn:
            db.cursor.execute("UPDATE users SET sub_end = ?, notified_expiry = 0 WHERE user_id = ?", (expire, user_id))
            
        # Обновляем Marzban напрямую
        username = f"user_{user_id}"
        res = marzban._request_with_retry(requests.get, f"/api/user/{username}")
        if res.status_code == 200:
            user_data = res.json()
            user_data['expire'] = expire
            user_data['status'] = "active"
            marzban._request_with_retry(
                requests.put, f"/api/user/{username}",
                json=user_data
            )
        else:
            payload = {
                "username": username,
                "proxies": {"vless": {"flow": "xtls-rprx-vision"}},
                "inbounds": {"vless": [cluster]} if cluster else {},
                "expire": expire,
                "data_limit": 0,
                "status": "active"
            }
            marzban._request_with_retry(requests.post, "/api/user", json=payload)
            
        bot.edit_message_text(
            f"⏳ Тестовая подписка на 1 минуту выдана пользователю {user_id}!",
            call.message.chat.id, call.message.message_id
        )
        try:
            bot.send_message(user_id, "⏳ Администратор выдал вам тестовую подписку на 1 минуту!")
        except Exception:
            pass
        return

    # Обработка удаления подписки
    if days == -2:
        cluster = db.get_user_cluster(user_id)
        if cluster:
            db.change_cluster_users_count(cluster, -1)
            db.set_user_cluster(user_id, None)
            
        db.revoke_subscription(user_id)
        marzban.delete_user(user_id)
        bot.edit_message_text(
            f"❌ Подписка пользователя {user_id} полностью удалена, доступ закрыт.",
            call.message.chat.id, call.message.message_id
        )
        try:
            bot.send_message(user_id, "❌ Ваша подписка была аннулирована администратором.")
        except Exception:
            pass
        return

    cluster = db.get_user_cluster(user_id)
    if not cluster:
        cluster = db.allocate_cluster()
        if not cluster:
            bot.edit_message_text(
                "❌ Нет свободных мест на серверах для выдачи новой подписки.",
                call.message.chat.id, call.message.message_id
            )
            return
        db.set_user_cluster(user_id, cluster)
        db.change_cluster_users_count(cluster, 1)

    db.add_subscription(user_id, days)
    marzban.add_days(user_id, days, cluster_tag=cluster)

    label = "♾️ Навсегда" if days == 0 else keyboards.format_days(days)
    bot.edit_message_text(
        f"✅ Подписка ({label}) выдана пользователю {user_id}!",
        call.message.chat.id, call.message.message_id
    )
    # Уведомляем пользователя
    try:
        bot.send_message(user_id, f"🎁 Администратор выдал вам подписку: {label}!")
    except Exception:
        pass

# ==========================================
#          ПРОМОКОДЫ
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_promo_menu")
def admin_promo_menu_handler(call):
    """Открывает подменю управления промокодами"""
    if call.from_user.id != Config.ADMIN_ID: return
    bot.edit_message_text(
        "🎫 <b>Управление промокодами</b>\n\nВыберите действие:",
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.admin_promo_menu(),
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda c: c.data == "admin_create_promo")
def create_promo_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите название промокода (на англ):")
    bot.register_next_step_handler(msg, get_promo_name)

def get_promo_name(message):
    name = message.text.upper().strip()
    msg = bot.send_message(message.chat.id, f"Введите сумму для промокода {name}:")
    bot.register_next_step_handler(msg, lambda m: get_promo_amount(m, name))

def get_promo_amount(message, name):
    try:
        amount = int(message.text)
        msg = bot.send_message(message.chat.id, f"Введите количество активаций для {name}:")
        bot.register_next_step_handler(msg, lambda m: create_promo_final(m, name, amount))
    except ValueError:
        bot.send_message(message.chat.id, "❌ Сумма должна быть числом.")

def create_promo_final(message, name, amount):
    try:
        activations = int(message.text)
        db.create_promo(name, amount, activations)
        bot.send_message(message.chat.id, f"✅ Промокод `{name}` на {amount}₽ ({activations} активаций) создан!", parse_mode="Markdown")
    except Exception as e:
        bot.send_message(message.chat.id, f"❌ Ошибка: {e}")

@bot.callback_query_handler(func=lambda c: c.data == "admin_list_promo")
def list_promos(call):
    if call.from_user.id != Config.ADMIN_ID: return
    try:
        promos = db.get_promos()
        if not promos:
            bot.send_message(call.message.chat.id, "Список пуст.")
            return
            
        text = "📋 **Активные промокоды:**\n\n"
        for p in promos:
            # p[0]=code, p[1]=amount, p[2]=activations_left
            text += f"🎫 `{p[0]}` — {p[1]}₽ (Осталось активаций: {p[2]})\n"
        
        bot.send_message(call.message.chat.id, text, parse_mode="Markdown")
    except Exception as e:
        bot.send_message(call.message.chat.id, f"Ошибка: {e}")

@bot.callback_query_handler(func=lambda c: c.data == "admin_del_one_promo")
def delete_promo_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите код для удаления:")
    bot.register_next_step_handler(msg, delete_promo_final)

def delete_promo_final(message):
    code = message.text.upper().strip()
    db.delete_promo_by_name(code)
    bot.send_message(message.chat.id, f"🗑 Промокод {code} удален (если существовал).")

import telebot.apihelper

@bot.callback_query_handler(func=lambda c: c.data == "admin_clear_inactive_promos")
def clear_inactive_promos(call):
    if call.from_user.id != Config.ADMIN_ID: return
    deleted_count = db.delete_inactive_promos()
    bot.answer_callback_query(call.id, f"✅ Удалено {deleted_count} неактивных промокодов", show_alert=True)
    try:
        admin_promo_menu_handler(call)
    except telebot.apihelper.ApiTelegramException as e:
        if "message is not modified" not in str(e):
            raise

# ==========================================
#          ТАРИФЫ (УПРАВЛЕНИЕ)
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_tariffs")
def admin_tariffs_list(call):
    """Показывает текущие тарифы с кнопками управления"""
    if call.from_user.id != Config.ADMIN_ID: return
    
    tariffs = db.get_tariffs()
    lines = ["📋 <b>Текущие тарифы:</b>\n"]
    for days, price in tariffs.items():
        label = keyboards.format_days(days)
        lines.append(f"  ⚡ {label} — <b>{price}₽</b>")
    
    text = "\n".join(lines)
    
    kb = types.InlineKeyboardMarkup(row_width=2)
    kb.add(
        types.InlineKeyboardButton("➕ Добавить", callback_data="admin_add_tariff"),
        types.InlineKeyboardButton("❌ Удалить", callback_data="admin_del_tariff")
    )
    kb.add(types.InlineKeyboardButton("🔙 Назад", callback_data="admin_back"))
    
    try:
        bot.edit_message_text(text, call.message.chat.id, call.message.message_id, 
                             reply_markup=kb, parse_mode="HTML")
    except Exception:
        bot.send_message(call.message.chat.id, text, reply_markup=kb, parse_mode="HTML")

@bot.callback_query_handler(func=lambda c: c.data == "admin_add_tariff")
def admin_add_tariff_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(call.message.chat.id, "Введите количество дней для нового тарифа:")
    bot.register_next_step_handler(msg, process_add_tariff_days)

def process_add_tariff_days(message):
    if not message.text or not message.text.isdigit() or int(message.text) <= 0:
        bot.send_message(message.chat.id, "❌ Ошибка! Введите положительное число дней.")
        return
    days = int(message.text)
    msg = bot.send_message(message.chat.id, f"Теперь введите цену для тарифа на {days} дней (в рублях):")
    bot.register_next_step_handler(msg, process_add_tariff_price, days)

def process_add_tariff_price(message, days):
    if not message.text or not message.text.isdigit() or int(message.text) <= 0:
        bot.send_message(message.chat.id, "❌ Ошибка! Цена должна быть положительным числом.")
        return
    price = int(message.text)
    db.add_tariff(days, price)
    
    label = keyboards.format_days(days)
    bot.send_message(message.chat.id, f"✅ Тариф добавлен: {label} — {price}₽")

@bot.callback_query_handler(func=lambda c: c.data == "admin_del_tariff")
def admin_del_tariff_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    
    tariffs = db.get_tariffs()
    if not tariffs:
        bot.answer_callback_query(call.id, "Нет тарифов для удаления!", show_alert=True)
        return
    
    kb = types.InlineKeyboardMarkup(row_width=1)
    for days, price in tariffs.items():
        label = keyboards.format_days(days)
        kb.add(types.InlineKeyboardButton(f"❌ {label} — {price}₽", callback_data=f"admin_rm_tariff_{days}"))
    kb.add(types.InlineKeyboardButton("🔙 Назад", callback_data="admin_tariffs"))
    
    bot.edit_message_text("Выберите тариф для удаления:", call.message.chat.id, 
                         call.message.message_id, reply_markup=kb)

@bot.callback_query_handler(func=lambda c: c.data.startswith("admin_rm_tariff_"))
def admin_rm_tariff(call):
    if call.from_user.id != Config.ADMIN_ID: return
    days = int(call.data.split("_")[3])
    
    if db.remove_tariff(days):
        label = keyboards.format_days(days)
        bot.answer_callback_query(call.id, f"✅ Тариф '{label}' удалён!", show_alert=True)
    else:
        bot.answer_callback_query(call.id, "❌ Тариф не найден!", show_alert=True)
    
    # Обновляем список тарифов
    admin_tariffs_list(call)

# ==========================================
#          ДАШБОРД (СТАТИСТИКА)
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_dashboard")
def admin_dashboard(call):
    if call.from_user.id != Config.ADMIN_ID: return
    
    s = db.get_stats()
    
    text = (
        "📊 <b>Дашборд</b>\n"
        "━━━━━━━━━━━━━━━━━━━\n\n"
        
        "👥 <b>Пользователи</b>\n"
        f"   Всего: <b>{s['total_users']}</b>\n"
        f"   Новых сегодня: <b>{s['new_today']}</b>\n"
        f"   Активных подписок: <b>{s['active_subs']}</b>\n\n"
        
        "💰 <b>Выручка</b>\n"
        f"   Сегодня: <b>{s['revenue_today']:.0f}₽</b>\n"
        f"   За 7 дней: <b>{s['revenue_week']:.0f}₽</b>\n"
        f"   За 30 дней: <b>{s['revenue_month']:.0f}₽</b>\n"
        f"   За всё время: <b>{s['revenue_total']:.0f}₽</b>\n\n"
        
        "📦 <b>Прочее</b>\n"
        f"   Оплаченных счетов: <b>{s['total_invoices']}</b>\n"
        f"   Активных промокодов: <b>{s['active_promos']}</b>\n"
        "━━━━━━━━━━━━━━━━━━━\n\n"
        
        "🤝 <b>Рефералы</b>\n"
        f"   Всего приглашено: <b>{s['total_referrals']}</b>\n"
        f"   Выплачено: <b>{s['total_ref_payouts']}₽</b>\n"
    )
    
    # Топ рефереров
    if s['top_referrers']:
        text += "\n🏆 <b>Топ рефереров:</b>\n"
        for i, ref in enumerate(s['top_referrers'], 1):
            text += f"   {i}. ID <code>{ref['user_id']}</code> — {ref['count']} чел., {ref['earned']}₽\n"
    
    text += "━━━━━━━━━━━━━━━━━━━"
    
    kb = types.InlineKeyboardMarkup()
    kb.add(types.InlineKeyboardButton("🔄 Обновить", callback_data="admin_dashboard"))
    kb.add(types.InlineKeyboardButton("🔙 Назад в админку", callback_data="admin_back"))
    
    try:
        bot.edit_message_text(text, call.message.chat.id, call.message.message_id, 
                             reply_markup=kb, parse_mode="HTML")
    except Exception:
        bot.send_message(call.message.chat.id, text, reply_markup=kb, parse_mode="HTML")



# ==========================================
#          КЛАСТЕРЫ
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_cluster_menu")
def admin_cluster_menu_handler(call):
    if call.from_user.id != Config.ADMIN_ID: return
    mode = db.get_setting("cluster_mode", "auto")
    sales_cluster = db.get_setting("sales_cluster", "")
    try:
        bot.edit_message_text(
            "🌐 <b>Кластеры (Inbounds)</b>\n\nУправляйте серверами и распределением пользователей:",
            call.message.chat.id, call.message.message_id,
            reply_markup=keyboards.admin_cluster_menu(mode, sales_cluster),
            parse_mode="HTML"
        )
    except Exception as e:
        if "message is not modified" not in str(e).lower():
            raise

@bot.callback_query_handler(func=lambda c: c.data == "admin_cluster_toggle_mode")
def admin_cluster_toggle_mode(call):
    if call.from_user.id != Config.ADMIN_ID: return
    current_mode = db.get_setting("cluster_mode", "auto")
    new_mode = "manual" if current_mode == "auto" else "auto"
    db.set_setting("cluster_mode", new_mode)
    
    bot.answer_callback_query(call.id, f"Режим изменен на {new_mode}")
    admin_cluster_menu_handler(call)

@bot.callback_query_handler(func=lambda c: c.data == "admin_cluster_select_sales")
def admin_cluster_select_sales(call):
    if call.from_user.id != Config.ADMIN_ID: return
    clusters = db.get_all_clusters()
    bot.edit_message_text(
        "🎯 Выберите кластер, в который будут попадать новые пользователи (даже сверх лимита):",
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.admin_cluster_select_menu(clusters),
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda c: c.data.startswith("admin_cluster_setsales_"))
def admin_cluster_setsales(call):
    if call.from_user.id != Config.ADMIN_ID: return
    tag = call.data.split("admin_cluster_setsales_")[1]
    
    if tag == "NONE":
        db.set_setting("sales_cluster", "")
        bot.answer_callback_query(call.id, "Выбор очищен")
    else:
        db.set_setting("sales_cluster", tag)
        bot.answer_callback_query(call.id, f"Кластер продаж: {tag}")
        
    admin_cluster_menu_handler(call)

@bot.callback_query_handler(func=lambda c: c.data == "admin_cluster_list")
def admin_cluster_list(call):
    if call.from_user.id != Config.ADMIN_ID: return
    clusters = db.get_all_clusters()
    if not clusters:
        bot.answer_callback_query(call.id, "Кластеры не найдены. Нажмите 'Обновить из Marzban'", show_alert=True)
        return
        
    bot.edit_message_text(
        "📋 <b>Список кластеров:</b>",
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.admin_cluster_list_menu(clusters),
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda c: c.data == "admin_cluster_sync")
def admin_cluster_sync(call):
    if call.from_user.id != Config.ADMIN_ID: return
    bot.answer_callback_query(call.id, "🔄 Синхронизация с Marzban...")
    inbounds = marzban.get_vless_inbounds()
    added = 0
    for tag in inbounds:
        if db.get_cluster(tag) is None:
            db.add_cluster(tag, max_users=10, is_active=0)
            added += 1
            
    bot.answer_callback_query(call.id, f"Синхронизировано. Добавлено: {added}")
    admin_cluster_menu_handler(call)

@bot.callback_query_handler(func=lambda c: c.data.startswith("admin_cluster_edit_"))
def admin_cluster_edit(call):
    if call.from_user.id != Config.ADMIN_ID: return
    tag = call.data.split("admin_cluster_edit_")[1]
    cluster = db.get_cluster(tag)
    if not cluster:
        bot.answer_callback_query(call.id, "Кластер не найден")
        return
        
    status = "🟢 Включен" if cluster['is_active'] else "🔴 Выключен"
    text = (
        f"⚙️ <b>Настройка кластера: {tag}</b>\n\n"
        f"Статус: {status}\n"
        f"Пользователи: {cluster['users_count']} / {cluster['max_users']}\n"
    )
    bot.edit_message_text(
        text,
        call.message.chat.id, call.message.message_id,
        reply_markup=keyboards.admin_cluster_edit_menu(cluster),
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda c: c.data.startswith("admin_cluster_toggle_") and not c.data.startswith("admin_cluster_toggle_mode"))
def admin_cluster_toggle(call):
    if call.from_user.id != Config.ADMIN_ID: return
    tag = call.data.split("admin_cluster_toggle_")[1]
    cluster = db.get_cluster(tag)
    if cluster:
        new_status = 0 if cluster['is_active'] else 1
        db.update_cluster(tag, cluster['max_users'], new_status)
        bot.answer_callback_query(call.id, "Статус изменен")
        # Перерисовываем меню
        call.data = f"admin_cluster_edit_{tag}"
        admin_cluster_edit(call)

@bot.callback_query_handler(func=lambda c: c.data.startswith("admin_cluster_setlimit_"))
def admin_cluster_setlimit_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    tag = call.data.split("admin_cluster_setlimit_")[1]
    
    msg = bot.send_message(call.message.chat.id, f"Введите новый лимит пользователей для кластера {tag}:", reply_markup=keyboards.cancel_admin_action())
    bot.register_next_step_handler(msg, process_cluster_limit, tag)

def process_cluster_limit(message, tag):
    if message.from_user.id != Config.ADMIN_ID: return
    if message.text == "🔙 Отмена": return
    
    if not message.text.isdigit() or int(message.text) < 0:
        bot.send_message(message.chat.id, "❌ Пожалуйста, введите положительное число.")
        return
        
    limit = int(message.text)
    cluster = db.get_cluster(tag)
    if cluster:
        db.update_cluster(tag, limit, cluster['is_active'])
        bot.send_message(message.chat.id, f"✅ Лимит для {tag} установлен на {limit}.")
    else:
        bot.send_message(message.chat.id, "❌ Кластер не найден.")


# ==========================================
#          РАССЫЛКА
# ==========================================
TEMP_BROADCAST = {}

@bot.callback_query_handler(func=lambda c: c.data == "admin_broadcast")
def broadcast_start(call):
    if call.from_user.id != Config.ADMIN_ID: return
    msg = bot.send_message(
        call.message.chat.id, 
        "Отправьте текст или фото (с текстом или без) для рассылки всем пользователям:", 
        reply_markup=keyboards.cancel_admin_action()
    )
    bot.register_next_step_handler(msg, process_broadcast_preview)

def process_broadcast_preview(message):
    admin_id = message.from_user.id
    if admin_id != Config.ADMIN_ID: return
    
    # Сохраняем сообщение
    TEMP_BROADCAST[admin_id] = message
    
    # Показываем предпросмотр
    bot.send_message(message.chat.id, "👀 **ПРЕДПРОСМОТР РАССЫЛКИ:**", parse_mode="Markdown")
    
    if message.content_type == 'text':
        bot.send_message(message.chat.id, message.text, reply_markup=keyboards.confirm_broadcast_menu())
    elif message.content_type == 'photo':
        bot.send_photo(message.chat.id, message.photo[-1].file_id, caption=message.caption, reply_markup=keyboards.confirm_broadcast_menu())
    else:
        bot.send_message(message.chat.id, "❌ Поддерживается только текст или фото.", reply_markup=keyboards.cancel_admin_action())

@bot.callback_query_handler(func=lambda c: c.data == "admin_cancel")
def admin_cancel_action(call):
    if call.from_user.id != Config.ADMIN_ID: return
    bot.clear_step_handler_by_chat_id(call.message.chat.id)
    try:
        bot.delete_message(call.message.chat.id, call.message.message_id)
    except Exception:
        pass
    bot.answer_callback_query(call.id, "Действие отменено.")

@bot.callback_query_handler(func=lambda c: c.data == "admin_confirm_broadcast")
def admin_confirm_broadcast(call):
    admin_id = call.from_user.id
    if admin_id != Config.ADMIN_ID: return
    
    saved_msg = TEMP_BROADCAST.get(admin_id)
    if not saved_msg:
        bot.answer_callback_query(call.id, "❌ Ошибка: сообщение не найдено!", show_alert=True)
        return
        
    try:
        bot.edit_message_reply_markup(chat_id=call.message.chat.id, message_id=call.message.message_id, reply_markup=None)
    except Exception:
        pass
    
    bot.send_message(call.message.chat.id, "⏳ Начинаю рассылку...")
    
    users = db.get_all_users() 
    count = 0
    
    for user in users:
        try:
            if isinstance(user, dict): uid = user.get('user_id')
            elif isinstance(user, (list, tuple)): uid = user[0]
            else: uid = user
                
            if saved_msg.content_type == 'text':
                bot.send_message(uid, saved_msg.text)
            elif saved_msg.content_type == 'photo':
                bot.send_photo(uid, saved_msg.photo[-1].file_id, caption=saved_msg.caption)
            count += 1
        except Exception:
            pass # Юзер заблочил бота
            
    bot.send_message(call.message.chat.id, f"✅ Рассылка завершена. Доставлено: {count} чел.")
    TEMP_BROADCAST.pop(admin_id, None)

# ==========================================
#          ЭКСПОРТ БД В EXCEL
# ==========================================
@bot.callback_query_handler(func=lambda c: c.data == "admin_export_db")
def admin_export_db(call):
    if call.from_user.id != Config.ADMIN_ID:
        return
    
    bot.answer_callback_query(call.id, "⏳ Формирую файл...")
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Удаляем дефолтный лист
    
    # Получаем список всех таблиц в БД
    db.cursor.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")
    tables = [row[0] for row in db.cursor.fetchall()]
    
    for table_name in tables:
        ws = wb.create_sheet(title=table_name)
        
        # Заголовки столбцов
        db.cursor.execute(f"PRAGMA table_info({table_name})")
        columns = [col[1] for col in db.cursor.fetchall()]
        ws.append(columns)
        
        # Данные
        db.cursor.execute(f"SELECT * FROM {table_name}")
        for row in db.cursor.fetchall():
            ws.append(list(row))
        
        # Автоширина столбцов
        for col_idx, col_name in enumerate(columns, 1):
            max_len = len(str(col_name))
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 50)
    
    # Сохраняем во временный файл
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = os.path.join(tempfile.gettempdir(), f"marzban_shop_bot_db_{timestamp}.xlsx")
    wb.save(filepath)
    
    # Отправляем файлом
    with open(filepath, "rb") as f:
        bot.send_document(call.message.chat.id, f, caption=f"📥 Экспорт БД ({timestamp})\nТаблиц: {len(tables)}")
    
    # Удаляем временный файл
    os.remove(filepath)

@bot.callback_query_handler(func=lambda c: c.data == "admin_export_raw")
def admin_export_raw(call):
    """Отправляет сырой .db файл через безопасный Backup API"""
    if call.from_user.id != Config.ADMIN_ID:
        return
    
    bot.answer_callback_query(call.id, "⏳ Создаю безопасную копию базы...")
    
    import os
    import sqlite3
    import time
    
    db_path = Config.DB_PATH
    backup_path = db_path + ".backup"
    
    try:
        # 1. Открываем соединения
        src = sqlite3.connect(db_path, timeout=10)
        dst = sqlite3.connect(backup_path)
        
        # 2. Делаем безопасный бэкап
        src.backup(dst)
        
        # 3. КРИТИЧЕСКИ ВАЖНО: закрываем соединения!
        # Это сбрасывает все буферы из оперативки на диск и снимает локи с файла
        src.close()
        dst.close()
        
        # 4. Читаем и отправляем готовый файл
        with open(backup_path, "rb") as f:
            bot.send_document(
                call.message.chat.id, 
                f, 
                visible_file_name=f"marzban_shop_bot_backup_{int(time.time())}.db",
                caption="📥 Безопасный экспорт базы данных (включая WAL)"
            )
            
    except Exception as e:
        bot.send_message(call.message.chat.id, f"❌ Ошибка создания бэкапа: {e}")
        
    finally:
        # 5. Убираем за собой в любом случае (даже если отправка упала с ошибкой сети)
        if os.path.exists(backup_path):
            os.remove(backup_path)
